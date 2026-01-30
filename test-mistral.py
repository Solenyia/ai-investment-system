import os
import uuid
import time
from docx import Document
from datetime import datetime

# Import tvojich funkcií z main.py a LangChain správy
from main import importer, create_supervisor_agent
from langchain_core.messages import ToolMessage, AIMessage

def extract_decision_from_stream(node_update):
    """
    Prehľadáva update z grafu a hľadá finálne rozhodnutie.
    """
    messages = node_update.get("messages", [])
    if not messages:
        return None

    for m in messages:
        content = str(m.content).upper()
        
        # 1. Kontrola ToolMessage (výstup z make_decision)
        if isinstance(m, ToolMessage):
            if "BUY" in content: return "BUY"
            if "SELL" in content or "AVOID" in content: return "SELL"
            if "HOLD" in content: return "HOLD"

        # 2. Kontrola AIMessage (text od agenta)
        if isinstance(m, AIMessage):
            if "SUCCESSFULLY TRANSFERRED" in content:
                continue
            if "BUY" in content: return "BUY"
            if "SELL" in content or "AVOID" in content: return "SELL"
            if "HOLD" in content: return "HOLD"
            
    return None

def run_benchmark_suite(stocks=["AAPL", "TSLA", "NVDA"], iterations=2):
    """
    Spustí MAS analýzu s meraním času.
    """
    all_results = []
    doc = Document()
    doc.add_heading(f'MAS Investment Report - {datetime.now().strftime("%Y-%m-%d")}', 0)
    
    # Celkový štart benchmarku
    overall_start_time = time.time()
    
    print("Inicializujem Supervisor agenta (GPT-4o)...")
    supervisor = create_supervisor_agent()

    for stock in stocks:
        stock = stock.upper()
        print(f"\n{'='*40}")
        print(f" ANALÝZA AKCIE: {stock}")
        print(f"{'='*40}")
        
        if f"{stock}_historical_data.csv" not in os.listdir():
            print(f"Sťahujem dáta pre {stock}...")
            importer(stock)

        for i in range(1, iterations + 1):
            print(f"\n>>> TEST {i}/{iterations} pre {stock} prebieha...")
            
            # Štart času pre jeden konkrétny test
            test_start_time = time.time()
            
            thread_id = f"test-{stock.lower()}-{uuid.uuid4().hex[:4]}"
            user_query = f"Analyze {stock} stock. I need a full debate and final decision."
            
            final_decision_found = "NEURČITÉ"
            
            # Streamovanie grafu
            for chunk in supervisor.stream(
                {"messages": [{"role": "user", "content": user_query}]},
                config={"configurable": {"thread_id": thread_id}}
            ):
                for node_name, node_update in chunk.items():
                    decision = extract_decision_from_stream(node_update)
                    if decision:
                        final_decision_found = decision

            # Koniec času pre test
            test_duration = time.time() - test_start_time
            
            # Uloženie výsledkov
            all_results.append({
                "akcia": stock,
                "test": i,
                "vysledok": final_decision_found,
                "cas": test_duration
            })
            
            # Zápis do Wordu
            doc.add_heading(f'{stock} - Test {i}', level=1)
            doc.add_paragraph(f"Finálne rozhodnutie: {final_decision_found}")
            doc.add_paragraph(f"Trvanie testu: {test_duration:.2f} sekúnd")
            
            print(f"   HOTOVO: {final_decision_found} (Čas: {test_duration:.2f}s)")

    # Celkový koniec benchmarku
    total_duration = time.time() - overall_start_time

    # ZÁVEREČNÁ TABUĽKA
    print("\n\n" + "="*50)
    print(f"{'VÝSLEDNÝ STAV VYHODNOTENIA (MAS)':^50}")
    print("="*50)
    print(f"{'AKCIA':<15} | {'TEST':<8} | {'ROZHODNUTIE':<15} | {'TRVANIE':<10}")
    print("-" * 75)
    
    for res in all_results:
        print(f"{res['akcia']:<15} | {res['test']:<8} | {res['vysledok']:<15} | {res['cas']:>6.2f} s")
    
    print("-" * 75)
    print(f"{'CELKOVÝ ČAS BENCHMARKU:':<43} {total_duration/60:>15.2f} min")
    print("="*50)

    # Uloženie reportu
    report_name = f"MAS_Benchmark_Report_{datetime.now().strftime('%H%M%S')}.docx"
    doc.save(report_name)
    print(f"\nKompletný report uložený v: {report_name}")

if __name__ == "__main__":
    # 3 akcie po 2 testy
    zoznam_akcii = ["AAPL", "TSLA", "NVDA"]
    run_benchmark_suite(stocks=zoznam_akcii, iterations=2)